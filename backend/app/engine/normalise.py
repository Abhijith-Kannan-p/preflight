"""Turn pasted text, Markdown, DOCX, PDF or an HTML email into a NormalisedDoc of text blocks
with offsets. Images and scanned PDFs become a VisualUpload, read later by OCR."""

import io
import re
from html.parser import HTMLParser
from pathlib import PurePosixPath
from typing import ClassVar

from app.engine.schema import Block, ImageRef, NormalisedDoc
from app.visual.intake import (
    IMAGE_EXTENSIONS,
    VisualIntakeError,
    VisualUpload,
    is_scanned,
    load_image,
    rasterise_pdf,
)

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_CHARS = 20_000

SUPPORTED_EXTENSIONS = {
    ".txt": "text",
    ".md": "markdown",
    ".docx": "docx",
    ".pdf": "pdf",
    ".html": "html",
    ".htm": "html",
    **{ext: "image" for ext in IMAGE_EXTENSIONS},
}


class IntakeError(ValueError):
    """Raised with a user-facing message when content cannot be read."""


def _blocks_from_paragraphs(
    paragraphs: list[str], source_format: str, allow_empty: bool = False
) -> NormalisedDoc:
    cleaned = [p.strip() for p in paragraphs if p and p.strip()]
    if not cleaned and not allow_empty:
        raise IntakeError("No readable text was found. Paste text or upload a text-based file.")
    total = sum(len(p) for p in cleaned) + 2 * (len(cleaned) - 1)
    if total > MAX_CHARS:
        raise IntakeError(f"Content is {total:,} characters; the limit is {MAX_CHARS:,}.")
    blocks: list[Block] = []
    offset = 0
    for i, paragraph in enumerate(cleaned):
        blocks.append(Block(id=i, text=paragraph, start_offset=offset))
        offset += len(paragraph) + 2
    return NormalisedDoc(blocks=blocks, source_format=source_format)


def _split_paragraphs(text: str) -> list[str]:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    return re.split(r"\n\s*\n", text)


_MD_RULES = [
    (re.compile(r"^\s{0,3}#{1,6}\s+", re.M), ""),
    (re.compile(r"^\s*[-*+]\s+", re.M), ""),
    (re.compile(r"!\[([^\]]*)\]\([^)]*\)"), r"\1"),
    (re.compile(r"\[([^\]]+)\]\([^)]*\)"), r"\1"),
    (re.compile(r"(\*\*|__)(.+?)\1"), r"\2"),
    (re.compile(r"(?<![\w*])\*(?!\s)(.+?)(?<!\s)\*(?![\w*])"), r"\1"),
    (re.compile(r"`([^`]+)`"), r"\1"),
]


def strip_markdown(text: str) -> str:
    for pattern, replacement in _MD_RULES:
        text = pattern.sub(replacement, text)
    return text


def from_text(text: str, markdown: bool = False) -> NormalisedDoc:
    if markdown:
        text = strip_markdown(text)
    return _blocks_from_paragraphs(_split_paragraphs(text), "markdown" if markdown else "text")


def from_docx(data: bytes) -> NormalisedDoc:
    from docx import Document

    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise IntakeError("This DOCX file could not be opened. Is it a valid Word file?") from exc
    paragraphs = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            paragraphs.append(" | ".join(cell.text.strip() for cell in row.cells))
    images = _docx_images(document)
    doc = _blocks_from_paragraphs(paragraphs, "docx", allow_empty=bool(images))
    doc.images = images
    return doc


def _docx_images(document: object) -> list[ImageRef]:
    """Inline pictures with their alt text (Word stores it as the drawing's `descr`)."""
    refs: list[ImageRef] = []
    for shape in getattr(document, "inline_shapes", []):
        doc_pr = shape._inline.docPr
        alt = doc_pr.get("descr") or doc_pr.get("title")
        refs.append(ImageRef(name=doc_pr.get("name") or "picture", alt=alt, source="docx"))
    return refs


def from_pdf(data: bytes, name: str = "document.pdf") -> NormalisedDoc | VisualUpload:
    """A PDF with a text layer becomes text blocks; a scan becomes page images for OCR."""
    import fitz

    try:
        document = fitz.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise IntakeError("This PDF could not be opened. Is it a valid PDF file?") from exc
    paragraphs: list[str] = []
    with document:
        if document.needs_pass:
            raise IntakeError("This PDF is password-protected.")
        if is_scanned(document):
            try:
                return rasterise_pdf(document, name)
            except VisualIntakeError as exc:
                raise IntakeError(str(exc)) from exc
        for page in document:
            for block in page.get_text("blocks", sort=True):
                if block[6] == 0:  # text block, not image
                    paragraphs.append(" ".join(block[4].split()))
    return _blocks_from_paragraphs(paragraphs, "pdf")


class _EmailParser(HTMLParser):
    """Text by block element, <img> tags with their alt attribute, and CSS colours."""

    BLOCK: ClassVar[set[str]] = {
        "p",
        "div",
        "td",
        "th",
        "li",
        "h1",
        "h2",
        "h3",
        "h4",
        "h5",
        "h6",
        "tr",
        "br",
    }
    SKIP: ClassVar[set[str]] = {"script", "style", "head", "title"}
    COLOUR = re.compile(
        r"(?:^|[;\s])(?:background-color|background|color)\s*:\s*(#[0-9a-f]{3,6})\b", re.I
    )

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.paragraphs: list[str] = [""]
        self.images: list[ImageRef] = []
        self.colours: list[str] = []
        self._skipping = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        values = dict(attrs)
        if tag in self.SKIP:
            self._skipping += 1
        if tag in self.BLOCK:
            self.paragraphs.append("")
        style = values.get("style") or ""
        self.colours += [m.group(1) for m in self.COLOUR.finditer(style)]
        for key in ("bgcolor", "color"):
            value = values.get(key)
            if value and re.fullmatch(r"#[0-9a-fA-F]{3,6}", value):
                self.colours.append(value)
        if tag == "img":
            src = values.get("src") or "image"
            name = PurePosixPath(src.split("?")[0]).name or "image"
            self.images.append(ImageRef(name=name, alt=values.get("alt"), source="html"))

    def handle_endtag(self, tag: str) -> None:
        if tag in self.SKIP and self._skipping:
            self._skipping -= 1
        if tag in self.BLOCK:
            self.paragraphs.append("")

    def handle_data(self, data: str) -> None:
        if not self._skipping:
            self.paragraphs[-1] += " " + " ".join(data.split())


def from_html(data: bytes) -> NormalisedDoc:
    """An HTML email: its text, its images' alt text and the colours its CSS uses. Remote
    images are not downloaded (upload an image itself for a visual review)."""
    try:
        markup = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        markup = data.decode("latin-1")
    parser = _EmailParser()
    parser.feed(markup)
    paragraphs = [" ".join(p.split()) for p in parser.paragraphs]
    doc = _blocks_from_paragraphs(paragraphs, "html", allow_empty=bool(parser.images))
    doc.images, doc.colours = parser.images, parser.colours
    return doc


def from_upload(filename: str, data: bytes) -> NormalisedDoc | VisualUpload:
    if len(data) > MAX_FILE_BYTES:
        raise IntakeError("Files must be 5 MB or smaller.")
    extension = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    kind = SUPPORTED_EXTENSIONS.get(extension)
    if kind is None:
        raise IntakeError(
            "Unsupported file type. Upload DOCX, PDF, TXT, MD, HTML or a PNG, JPG or WebP image."
        )
    if kind == "docx":
        return from_docx(data)
    if kind == "pdf":
        return from_pdf(data, filename)
    if kind == "html":
        return from_html(data)
    if kind == "image":
        try:
            return load_image(filename, data)
        except VisualIntakeError as exc:
            raise IntakeError(str(exc)) from exc
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise IntakeError("Text files must be UTF-8 encoded.") from exc
    return from_text(text, markdown=kind == "markdown")
