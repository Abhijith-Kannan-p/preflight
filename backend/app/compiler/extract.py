"""Read a brand book into pages, and the deterministic extractors for its structured parts.

Pages keep their numbers (so every candidate cites `source_doc#pN`). Tables are read with
PyMuPDF's table finder and classified by their header row. Free text is scanned with a few
transparent patterns: numeric limits ("at most one exclamation mark per 100 words"), required
disclaimers ('must include "Subject to eligibility."') and quoted claim phrases under claim
headings. Everything else is prose, left to the LLM step.
"""

import io
import re
from dataclasses import dataclass, field
from typing import Any

from app.compiler.schema import Candidate, validate

PAGE_FOOTER = re.compile(r"^.*\bpage \d+\s*$", re.I)
QUOTED = re.compile(r"[\"“]([^\"”]{1,60})[\"”]")
NUMBER_WORDS = {"one": 1, "two": 2, "three": 3, "four": 4, "five": 5, "ten": 10}


class CompileInputError(ValueError):
    """User-facing: the document could not be read as a brand book."""


@dataclass
class Page:
    number: int
    title: str
    text: str
    tables: list[list[list[str]]] = field(default_factory=list)


def read_pdf(data: bytes) -> list[Page]:
    import fitz

    try:
        document = fitz.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise CompileInputError("This PDF could not be opened.") from exc
    if document.needs_pass:
        raise CompileInputError("This PDF is password-protected.")
    pages: list[Page] = []
    for index, page in enumerate(document):
        found = page.find_tables()
        tables = [[[(c or "").strip() for c in row] for row in t.extract()] for t in found.tables]
        cells = {c for t in tables for row in t for c in row if c}
        lines = [
            line.strip()
            for line in page.get_text("text").splitlines()
            if line.strip() and not PAGE_FOOTER.match(line.strip())
        ]
        title = lines[0] if lines else ""
        prose = [line for line in lines[1:] if line not in cells]
        pages.append(Page(index + 1, title, _join(prose), tables))
    if not any(p.text or p.tables for p in pages):
        raise CompileInputError("No text was found. Upload a brand book with a text layer.")
    return pages


def read_docx(data: bytes) -> list[Page]:
    """DOCX has no fixed pages: headings start sections, numbered in order as 'pages'."""
    from docx import Document

    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise CompileInputError("This DOCX file could not be opened.") from exc
    pages: list[Page] = []
    for paragraph in document.paragraphs:
        style = (paragraph.style.name if paragraph.style is not None else "").lower()
        if style.startswith("heading") or not pages:
            pages.append(Page(len(pages) + 1, paragraph.text.strip(), ""))
        elif paragraph.text.strip():
            pages[-1].text = (pages[-1].text + " " + paragraph.text.strip()).strip()
    for table in document.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        (pages[-1] if pages else Page(1, "", "")).tables.append(rows)
    return pages


def _join(lines: list[str]) -> str:
    """Re-flow wrapped lines into paragraphs."""
    text = " ".join(lines)
    return re.sub(r"\s+", " ", text).strip()


def sentences(text: str) -> list[str]:
    parts = re.split(r"(?:(?<=[.!?])|(?<=[.!?][\"”]))\s+(?=[A-Z\"“])", text)
    return [s.strip() for s in parts if s.strip()]


# --- tables ----------------------------------------------------------------------------


def _number(cell: str) -> int | None:
    match = re.search(r"\d+", cell)
    return int(match.group()) if match else None


def _split_variants(cell: str) -> list[str]:
    return [v.strip() for v in re.split(r",|;|\|", cell) if v.strip() and v.strip() != "-"]


def from_table(rows: list[list[str]], page: int) -> list[Candidate]:
    rows = [[" ".join(c.split()) for c in row] for row in rows]  # wrapped cells
    if len(rows) < 2:
        return []
    header = [h.lower() for h in rows[0]]
    body = [r for r in rows[1:] if any(r)]
    found: list[Candidate] = []

    def add(kind: str, payload: dict[str, Any], row: list[str]) -> None:
        valid = validate(kind, payload)
        if valid is not None:
            found.append(
                Candidate(kind=kind, payload=valid, source_page=page, excerpt=" | ".join(row))
            )

    if any("never write" in h or "variant" in h for h in header):
        for row in body:
            add(
                "term",
                {
                    "term": row[0],
                    "variants": _split_variants(row[1]),
                    "note": row[2] if len(row) > 2 else "",
                },
                row,
            )
    elif any(h == "status" for h in header) and any("word" in h for h in header):
        for row in body:
            add(
                "word",
                {
                    "term": row[0],
                    "status": row[1].lower(),
                    "replacement": row[2] if len(row) > 2 else None,
                },
                row,
            )
    elif any("hex" in h for h in header):
        for row in body:
            add("colour", {"name": row[0], "hex": row[1]}, row)
    elif any("channel" in h for h in header):
        for row in body:
            add(
                "channel",
                {
                    "channel": row[0],
                    "max_chars": _number(row[1]) if len(row) > 1 else None,
                    "max_hashtags": _number(row[2]) if len(row) > 2 else None,
                    "subject_max_chars": _number(row[3]) if len(row) > 3 else None,
                },
                row,
            )
    elif any("trait" in h for h in header) and any("target" in h for h in header):
        traits = {}
        for row in body:
            try:
                traits[row[0].strip().lower()] = float(row[1])
            except (ValueError, IndexError):
                continue
        add("voice", {"traits": traits}, [f"{k} {v}" for k, v in traits.items()])
    return found


# --- free text -------------------------------------------------------------------------

THRESHOLD_PATTERNS = [
    (r"at most (\w+) exclamation marks? per 100 words", "TON-03", "max_exclamations_per_100_words"),
    (r"capitals? longer than (\w+) letters", "TON-03", "max_caps_word_length"),
    (r"flesch reading ease score of (\d+)", "RDB-01", "min_flesch"),
    (r"sentences? to (\d+) words", "RDB-01", "max_sentence_words"),
]
CLAIM_HEADINGS = [
    (("superlative",), "CLM-01"),
    (("security", "safety"), "CLM-04"),
    (("promise", "guarantee"), "CLM-06"),
    (("free",), "CLM-08"),
]
DISCLAIMER = re.compile(r"must (?:include|carry)\s+[\"“]([^\"”]+)[\"”]", re.I)
MENTION = re.compile(
    r"mention of (?:the )?(?:[\"“]([^\"”]+)[\"”]|PayWise\s+([A-Z][\w ]+?)\s+must)", re.I
)


def _value(word: str) -> float | None:
    if word.isdigit():
        return float(word)
    return float(NUMBER_WORDS[word.lower()]) if word.lower() in NUMBER_WORDS else None


def from_text(page: Page) -> list[Candidate]:
    found: list[Candidate] = []
    lowered = page.text.lower()
    for pattern, rule, field_name in THRESHOLD_PATTERNS:
        match = re.search(pattern, lowered)
        value = _value(match.group(1)) if match else None
        if match and value is not None:
            payload = validate("threshold", {"rule": rule, "field": field_name, "value": value})
            if payload:
                found.append(
                    Candidate(
                        kind="threshold",
                        payload=payload,
                        source_page=page.number,
                        excerpt=_around(page.text, match.start()),
                    )
                )

    title = page.title.lower()
    for words, rule in CLAIM_HEADINGS:
        if any(w in title for w in words):
            first = sentences(page.text)[:1]
            for phrase in QUOTED.findall(first[0] if first else ""):
                if phrase.strip().endswith("."):
                    continue  # a quoted disclaimer, not a phrase to avoid
                payload = validate("pattern", {"rule": rule, "pattern": phrase.strip()})
                if payload:
                    found.append(
                        Candidate(
                            kind="pattern",
                            payload=payload,
                            source_page=page.number,
                            excerpt=first[0],
                        )
                    )

    for sentence in sentences(page.text):
        disclaimer = DISCLAIMER.search(sentence)
        if not disclaimer:
            continue
        mention = MENTION.search(sentence)
        triggers = []
        if mention:
            triggers.append((mention.group(1) or mention.group(2) or "").strip().lower())
        else:
            triggers += [q.lower() for q in QUOTED.findall(sentence[: disclaimer.start()])]
        payload = validate(
            "disclaimer",
            {
                "title": f"{disclaimer.group(1).rstrip('.')} disclaimer",
                "triggers": [t for t in triggers if t],
                "disclaimer": disclaimer.group(1).strip(),
            },
        )
        if payload:
            found.append(
                Candidate(
                    kind="disclaimer", payload=payload, source_page=page.number, excerpt=sentence
                )
            )
    return found


def _around(text: str, index: int, width: int = 160) -> str:
    start = max(0, text.rfind(".", 0, index) + 1)
    end = text.find(".", index)
    return text[start : end + 1 if end != -1 else start + width].strip()


def deterministic(pages: list[Page]) -> list[Candidate]:
    found: list[Candidate] = []
    for page in pages:
        for rows in page.tables:
            found += from_table(rows, page.number)
        found += from_text(page)
    return found
