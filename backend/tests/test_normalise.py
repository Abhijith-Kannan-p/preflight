import io

import pytest
from docx import Document

from app.config import get_settings
from app.engine.normalise import MAX_CHARS, IntakeError, from_text, from_upload


def test_text_blocks_keep_offsets_into_joined_text() -> None:
    doc = from_text("First para.\n\nSecond para\nwith a line break.\n\n\nThird.")
    assert [b.text for b in doc.blocks] == [
        "First para.",
        "Second para\nwith a line break.",
        "Third.",
    ]
    for block in doc.blocks:
        assert doc.text[block.start_offset : block.start_offset + len(block.text)] == block.text


def test_markdown_is_stripped_to_plain_text() -> None:
    doc = from_text("# Title\n\n**Bold** and [a link](https://x.y) and `code`.", markdown=True)
    assert [b.text for b in doc.blocks] == ["Title", "Bold and a link and code."]


def test_docx_paragraphs_become_blocks() -> None:
    buffer = io.BytesIO()
    document = Document()
    document.add_paragraph("Hello PayWise.")
    document.add_paragraph("")
    document.add_paragraph("Second paragraph.")
    document.save(buffer)
    doc = from_upload("draft.docx", buffer.getvalue())
    assert [b.text for b in doc.blocks] == ["Hello PayWise.", "Second paragraph."]
    assert doc.source_format == "docx"


def test_sample_pdf_and_docx_are_readable() -> None:
    files = get_settings().seed_dir / "files"
    pdf = from_upload("p.pdf", (files / "paywise_press_release.pdf").read_bytes())
    assert any("Project Falcon" in b.text for b in pdf.blocks)
    docx = from_upload("n.docx", (files / "paywise_october_newsletter.docx").read_bytes())
    assert docx.blocks[0].text.startswith("Subject:")


@pytest.mark.parametrize(
    ("filename", "data", "message"),
    [
        ("x.exe", b"MZ", "Unsupported file type"),
        ("x.pdf", b"not a pdf", "could not be opened"),
        ("x.docx", b"not a docx", "could not be opened"),
        ("x.txt", b"   \n\n  ", "No readable text"),
        ("x.txt", b"\xff\xfe\xfa", "UTF-8"),
        ("x.txt", b"a" * (5 * 1024 * 1024 + 1), "5 MB"),
    ],
    ids=["exe", "bad-pdf", "bad-docx", "empty", "not-utf8", "too-big"],
)
def test_bad_uploads_raise_friendly_errors(filename: str, data: bytes, message: str) -> None:
    with pytest.raises(IntakeError, match=message):
        from_upload(filename, data)


def test_text_over_limit_is_rejected() -> None:
    with pytest.raises(IntakeError, match="limit"):
        from_text("word " * (MAX_CHARS // 4))
